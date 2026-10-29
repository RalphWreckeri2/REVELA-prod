"""PSIC line matching and non-blocking registry sector-map maintenance."""

import logging
import json
import math
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


_LOGGER = logging.getLogger("revela.analytics.sectors")
_NON_ALPHANUMERIC = re.compile(r"[^a-z0-9]")
_SECTION_H = re.compile(r"^SECTION\s+H\.\s*", re.IGNORECASE)
_TEST_LINE = "sample ko"
_DISPLAY_SECTOR_DEFAULTS_PATH = Path(__file__).with_name("display_sector_defaults.json")
with _DISPLAY_SECTOR_DEFAULTS_PATH.open(encoding="utf-8") as _defaults_file:
    _DISPLAY_SECTOR_DEFAULTS = json.load(_defaults_file)


def normalize_sector_text(value):
    """Normalize master and registry descriptions using the approved rule."""
    return _NON_ALPHANUMERIC.sub("", str(value or "").lower())


def clean_major_type(value):
    cleaned = _SECTION_H.sub("", str(value or "").strip()).strip().upper()
    if cleaned == "FINANCIAL AND INSURANCE ACTIVITIE":
        return "FINANCIAL AND INSURANCE ACTIVITIES"
    return cleaned


def read_psic_workbook(path, source_version):
    """Read versioned PSIC rows, retaining duplicate normalized descriptions."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        header = [str(cell.value or "").strip().lower() for cell in next(sheet.iter_rows())]
        required = ("no.", "major type", "division", "group", "class", "line of business", "total")
        missing = [name for name in required if name not in header]
        if missing:
            raise ValueError(f"PSIC workbook is missing required columns: {', '.join(missing)}")
        indexes = {name: header.index(name) for name in required}
        rows = []
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            line = str(values[indexes["line of business"]] or "").strip()
            normalized_line = normalize_sector_text(line)
            if not normalized_line or normalized_line == _TEST_LINE:
                continue
            major_raw = str(values[indexes["major type"]] or "")
            major_clean = clean_major_type(major_raw)
            if not major_clean:
                raise ValueError(f"PSIC workbook row {row_number} has no Major Type.")
            total = values[indexes["total"]] or 0
            rows.append({
                "source_version": source_version,
                "source_row_number": int(values[indexes["no."]] or row_number - 1),
                "major_type_raw": major_raw,
                "major_type_clean": major_clean,
                "normalized_major_type": normalize_sector_text(major_clean),
                "division": values[indexes["division"]],
                "group": values[indexes["group"]],
                "class": values[indexes["class"]],
                "line_of_business": line,
                "normalized_line_of_business": normalized_line,
                "source_total": int(total),
            })
        return rows
    finally:
        workbook.close()


def display_sector_for_major(major_type):
    mapping = _DISPLAY_SECTOR_DEFAULTS.get(clean_major_type(major_type))
    return tuple(mapping) if mapping else None


def decompose_registry_activity(raw_value, psic_lines, display_sectors):
    """Return exact or longest-prefix PSIC matches, preserving real ambiguity."""
    normalized = normalize_sector_text(raw_value)
    if not normalized:
        return {"status": "unmatched", "normalized": normalized, "lines": [], "candidates": 0}

    by_normalized = {}
    for line in psic_lines:
        key = line["normalized_line_of_business"]
        by_normalized.setdefault(key, []).append(line)
    keys = sorted(by_normalized, key=len, reverse=True)

    position = 0
    matched = []
    ambiguous = False
    candidate_count = 0
    while position < len(normalized):
        candidates = [
            key for key in keys
            if normalized.startswith(key, position)
        ]
        if not candidates:
            return {
                "status": "unmatched", "normalized": normalized,
                "lines": matched, "candidates": 0,
            }
        longest = candidates[0]
        options = by_normalized[longest]
        candidate_count += len(options)
        if len(options) != 1:
            option_sectors = {
                display_sectors.get(option["normalized_major_type"], {}).get(
                    "display_sector_key"
                )
                for option in options
            }
            if len(option_sectors) == 1 and None not in option_sectors:
                options = [
                    min(
                        options,
                        key=lambda option: (
                            option.get("source_row_number", option.get("psic_line_id", 0)),
                            option.get("psic_line_id", 0),
                        ),
                    )
                ]
                matched.append(options[0])
            else:
                ambiguous = True
                matched.extend(options)
        else:
            matched.append(options[0])
        position += len(longest)

    if not matched:
        return {"status": "unmatched", "normalized": normalized, "lines": [], "candidates": 0}
    if ambiguous:
        return {
            "status": "ambiguous", "normalized": normalized,
            "lines": matched, "candidates": candidate_count,
        }

    primary = matched[0]
    sector = display_sectors.get(primary["normalized_major_type"])
    if not sector or not sector.get("display_sector_key"):
        return {
            "status": "unmatched", "normalized": normalized,
            "lines": matched, "candidates": 0,
        }
    return {
        "status": "matched" if len(matched) == 1 else "combo",
        "normalized": normalized,
        "lines": matched,
        "primary_sector_key": sector["display_sector_key"],
        "primary_sector_label": sector["display_sector_label"],
        "candidates": candidate_count,
    }


def _active_sector_data(cursor):
    cursor.execute("SELECT source_version FROM psic_active_version WHERE singleton_id = 1")
    active = cursor.fetchone()
    if not active:
        raise RuntimeError("No active PSIC master version is loaded.")
    source_version = active["source_version"] if isinstance(active, dict) else active[0]

    cursor.execute(
        """SELECT psic_line_id, source_row_number, normalized_line_of_business,
                  normalized_major_type, major_type_clean, line_of_business
           FROM psic_lines WHERE source_version = %s""",
        (source_version,),
    )
    line_rows = cursor.fetchall()
    cursor.execute(
        """SELECT normalized_major_type, display_sector_key, display_sector_label
           FROM display_sector
           WHERE is_active = 1"""
    )
    sector_rows = cursor.fetchall()
    display_sectors = {
        row["normalized_major_type"]: row for row in sector_rows
    }
    return source_version, line_rows, display_sectors


def sync_registry_sector_values(cursor, values):
    """Materialize classifications and queue ambiguous/unmatched values."""
    source_version, psic_lines, display_sectors = _active_sector_data(cursor)
    raw_values = set()
    for value in values:
        if value is None or (isinstance(value, float) and math.isnan(value)):
            continue
        raw_value = str(value).strip()
        if raw_value:
            raw_values.add(raw_value)
    raw_values = sorted(raw_values)
    outcomes = Counter()
    for raw_value in raw_values:
        result = decompose_registry_activity(raw_value, psic_lines, display_sectors)
        normalized = result["normalized"]
        primary_id = result["lines"][0]["psic_line_id"] if result["lines"] and result["status"] in ("matched", "combo") else None
        sector_key = result.get("primary_sector_key")
        sector_label = result.get("primary_sector_label")
        cursor.execute(
            """INSERT INTO registry_sector_map
                   (normalized_line_of_business, raw_line_of_business, source_version,
                    match_status, primary_psic_line_id, display_sector_key,
                    display_sector_label, candidate_count)
               VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
               ON DUPLICATE KEY UPDATE raw_line_of_business = VALUES(raw_line_of_business),
                   source_version = VALUES(source_version), match_status = VALUES(match_status),
                   primary_psic_line_id = VALUES(primary_psic_line_id),
                   display_sector_key = VALUES(display_sector_key),
                   display_sector_label = VALUES(display_sector_label),
                   candidate_count = VALUES(candidate_count)""",
            (
                normalized, raw_value, source_version, result["status"],
                primary_id, sector_key, sector_label, result["candidates"],
            ),
        )
        cursor.execute(
            "DELETE FROM registry_sector_map_secondary WHERE normalized_line_of_business = %s",
            (normalized,),
        )
        if result["status"] in ("matched", "combo"):
            for line in result["lines"][1:]:
                cursor.execute(
                    """INSERT IGNORE INTO registry_sector_map_secondary
                           (normalized_line_of_business, psic_line_id) VALUES (%s, %s)""",
                    (normalized, line["psic_line_id"]),
                )
            cursor.execute(
                """UPDATE registry_sector_review_queue
                   SET review_status = 'resolved', candidate_count = %s,
                       last_seen_at = CURRENT_TIMESTAMP
                   WHERE normalized_line_of_business = %s""",
                (result["candidates"], normalized),
            )
        else:
            cursor.execute(
                """INSERT INTO registry_sector_review_queue
                       (normalized_line_of_business, raw_line_of_business, review_status,
                        candidate_count, first_seen_at, last_seen_at)
                   VALUES (%s, %s, %s, %s, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
                   ON DUPLICATE KEY UPDATE raw_line_of_business = VALUES(raw_line_of_business),
                       review_status = VALUES(review_status),
                       candidate_count = VALUES(candidate_count),
                       last_seen_at = CURRENT_TIMESTAMP""",
                (normalized, raw_value, result["status"], result["candidates"]),
            )
        outcomes[result["status"]] += 1
    return outcomes


def sync_imported_registry_sector_values(connection, values):
    """Best-effort post-commit mapping sync; classification never blocks import."""
    cursor = None
    try:
        cursor = connection.cursor()
        sync_registry_sector_values(cursor, values)
        connection.commit()
    except Exception:
        connection.rollback()
        _LOGGER.exception(
            "Registry import completed, but sector-map synchronization failed; "
            "the import was not rolled back."
        )
    finally:
        if cursor is not None:
            cursor.close()


def seed_display_sector_rows(cursor, source_version, major_types):
    """Insert default Major Type mappings without overwriting staff edits."""
    rows = []
    for major_raw, major_clean, normalized_major in major_types:
        sector = display_sector_for_major(major_clean)
        key, label, order = sector if sector else (None, None, 99)
        rows.append((
            source_version, normalized_major, major_raw, major_clean,
            key, label, order,
        ))
    cursor.executemany(
        """INSERT INTO display_sector
               (normalized_major_type, source_version, major_type_raw,
                major_type_clean, display_sector_key, display_sector_label,
                sort_order, is_active)
           VALUES (%s, %s, %s, %s, %s, %s, %s, 1)
           ON DUPLICATE KEY UPDATE source_version = VALUES(source_version),
               major_type_raw = VALUES(major_type_raw),
               major_type_clean = VALUES(major_type_clean)""",
        [
            (
                normalized_major, source_version, major_raw, major_clean,
                key, label, order,
            )
            for source_version, normalized_major, major_raw, major_clean,
            key, label, order in rows
        ],
    )
